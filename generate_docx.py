import os
import sys

try:
    import docx
except ImportError:
    print("Installing python-docx...")
    os.system(f"{sys.executable} -m pip install python-docx")
    import docx

from docx import Document
from docx.shared import Twips, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# Measurements extracted from CV (2).docx (twips; 20 twips = 1 pt, 1440 = 1 in)
PAGE_W = 11906          # A4 width
PAGE_H = 16838          # A4 height
MARGIN = 864            # 0.6 in
CONTENT_RIGHT_TAB = 10181
DEFAULT_TAB = 720
BODY_SIZE = Pt(11)      # Normal / inherited (sz 22)
NAME_SIZE = Pt(24)      # sz 48
SECTION_SIZE = Pt(12)   # sz 24
SKILL_SIZE = Pt(10)     # sz 20
BULLET_SIZE = Pt(10)    # sz 20
TECH_SIZE = Pt(10)      # sz 20
LINK_COLOR = "0046B4"
FONT_NAME = "Times New Roman"


# ── XML Styling Helpers ───────────────────────────────────────────────────────

def set_rfonts(rPr, name=FONT_NAME):
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    rFonts.set(qn("w:ascii"), name)
    rFonts.set(qn("w:hAnsi"), name)
    rFonts.set(qn("w:eastAsia"), name)
    rFonts.set(qn("w:cs"), name)


def style_run(run, *, size=None, bold=None, italic=None, font=FONT_NAME):
    """Apply Times New Roman and optional size/weight. Omit size to inherit 11pt."""
    run.font.name = font
    rPr = run._element.get_or_add_rPr()
    set_rfonts(rPr, font)
    if size is not None:
        run.font.size = size
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    return run


def add_tab(paragraph):
    run = paragraph.add_run()
    run._r.append(OxmlElement("w:tab"))
    return run


def set_spacing(paragraph, *, before=0, after=0, line=240, line_rule="auto"):
    """Set w:spacing in twips to match the reference XML exactly."""
    pPr = paragraph._p.get_or_add_pPr()
    spacing = pPr.find(qn("w:spacing"))
    if spacing is None:
        spacing = OxmlElement("w:spacing")
        pPr.append(spacing)
    spacing.set(qn("w:before"), str(before))
    spacing.set(qn("w:after"), str(after))
    if line is not None:
        spacing.set(qn("w:line"), str(line))
        spacing.set(qn("w:lineRule"), line_rule)
    else:
        if qn("w:line") in spacing.attrib:
            del spacing.attrib[qn("w:line")]
        if qn("w:lineRule") in spacing.attrib:
            del spacing.attrib[qn("w:lineRule")]


def set_indent(paragraph, *, start=None, hanging=None):
    pPr = paragraph._p.get_or_add_pPr()
    ind = pPr.find(qn("w:ind"))
    if ind is None:
        ind = OxmlElement("w:ind")
        pPr.append(ind)
    if start is not None:
        ind.set(qn("w:start"), str(start))
    if hanging is not None:
        ind.set(qn("w:hanging"), str(hanging))


def set_right_tabs(paragraph):
    """Clear the default 0.5in tab and place a right tab at the content edge."""
    pPr = paragraph._p.get_or_add_pPr()
    existing = pPr.find(qn("w:tabs"))
    if existing is not None:
        pPr.remove(existing)
    tabs = OxmlElement("w:tabs")
    clear = OxmlElement("w:tab")
    clear.set(qn("w:val"), "clear")
    clear.set(qn("w:pos"), str(DEFAULT_TAB))
    right = OxmlElement("w:tab")
    right.set(qn("w:val"), "right")
    right.set(qn("w:pos"), str(CONTENT_RIGHT_TAB))
    right.set(qn("w:leader"), "none")
    tabs.append(clear)
    tabs.append(right)
    pPr.append(tabs)


def add_hyperlink(paragraph, text, url):
    """Clickable hyperlink: inherit 11pt TNR, #0046B4, single underline."""
    part = paragraph.part
    r_id = part.relate_to(url, docx.opc.constants.RELATIONSHIP_TYPE.HYPERLINK, is_external=True)

    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)

    new_run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")

    r_style = OxmlElement("w:rStyle")
    r_style.set(qn("w:val"), "Style")
    rPr.append(r_style)

    c = OxmlElement("w:color")
    c.set(qn("w:val"), LINK_COLOR)
    rPr.append(c)

    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rPr.append(u)

    new_run.append(rPr)
    t = OxmlElement("w:t")
    t.text = text
    new_run.append(t)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def add_bottom_border(paragraph):
    """Section-title rule: single, 0.75pt (sz 6), space 1, black."""
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "000000")
    pBdr.append(bottom)
    pPr.append(pBdr)


def set_keep_with_next(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    existing = pPr.find(qn("w:keepNext"))
    if existing is not None:
        pPr.remove(existing)
    kn = OxmlElement("w:keepNext")
    kn.set(qn("w:val"), "true")
    pPr.append(kn)


def set_document_defaults(doc):
    """Match CV (2).docx Normal + docDefaults: TNR 11pt, single spacing, no extra gap."""
    normal = doc.styles["Normal"]
    normal.font.name = FONT_NAME
    normal.font.size = BODY_SIZE
    rPr = normal.element.get_or_add_rPr()
    set_rfonts(rPr)
    color = rPr.find(qn("w:color"))
    if color is None:
        color = OxmlElement("w:color")
        rPr.append(color)
    color.set(qn("w:val"), "auto")

    pf = normal.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = 1.0

    # Document-level run defaults
    styles_el = doc.styles.element
    doc_defaults = styles_el.find(qn("w:docDefaults"))
    if doc_defaults is None:
        doc_defaults = OxmlElement("w:docDefaults")
        styles_el.insert(0, doc_defaults)
    rPrDefault = doc_defaults.find(qn("w:rPrDefault"))
    if rPrDefault is None:
        rPrDefault = OxmlElement("w:rPrDefault")
        doc_defaults.append(rPrDefault)
    def_rPr = rPrDefault.find(qn("w:rPr"))
    if def_rPr is None:
        def_rPr = OxmlElement("w:rPr")
        rPrDefault.append(def_rPr)
    set_rfonts(def_rPr)
    for tag, val in (("w:sz", "22"), ("w:szCs", "22")):
        el = def_rPr.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            def_rPr.append(el)
        el.set(qn("w:val"), val)

    # Default tab stop 720 twips (0.5 in), as in the reference
    settings = doc.settings.element
    dts = settings.find(qn("w:defaultTabStop"))
    if dts is None:
        dts = OxmlElement("w:defaultTabStop")
        settings.append(dts)
    dts.set(qn("w:val"), str(DEFAULT_TAB))


# ── Layout Elements ───────────────────────────────────────────────────────────

def add_section_header(doc, title):
    p = doc.add_paragraph()
    set_keep_with_next(p)
    add_bottom_border(p)
    set_spacing(p, before=280, after=80, line=240)
    run = p.add_run(title.upper())
    style_run(run, size=SECTION_SIZE, bold=True)
    return p


def add_bullet(doc, text):
    """Manual hanging bullet: start 360, hanging 216, after 40, 10pt, • + tab + text."""
    p = doc.add_paragraph()
    set_spacing(p, before=0, after=40, line=240)
    set_indent(p, start=360, hanging=216)

    run = p.add_run()
    style_run(run, size=BULLET_SIZE)
    t_bullet = OxmlElement("w:t")
    t_bullet.text = "•"
    run._r.append(t_bullet)
    run._r.append(OxmlElement("w:tab"))
    t_text = OxmlElement("w:t")
    t_text.text = text
    run._r.append(t_text)
    return p


def add_bullet_with_parts(doc, parts):
    """Manual hanging bullet with mixed text and hyperlinks.
    parts: list of tuples (text, url) where url can be None for plain text.
    """
    p = doc.add_paragraph()
    set_spacing(p, before=0, after=40, line=240)
    set_indent(p, start=360, hanging=216)

    run = p.add_run()
    style_run(run, size=BULLET_SIZE)
    t_bullet = OxmlElement("w:t")
    t_bullet.text = "•"
    run._r.append(t_bullet)
    run._r.append(OxmlElement("w:tab"))

    for text, url in parts:
        if url:
            add_hyperlink(p, text, url)
        else:
            r = p.add_run(text)
            style_run(r, size=BULLET_SIZE)
    return p


def add_experience_header(doc, company, dates, role, location, first=False):
    # Company (Bold 11pt) ... Dates (Bold 11pt, right-aligned)
    # First entry after the section heading has before=0; later entries keep 240.
    p1 = doc.add_paragraph()
    set_keep_with_next(p1)
    set_right_tabs(p1)
    set_spacing(p1, before=0 if first else 240, after=0, line=240)
    set_indent(p1, start=0)
    style_run(p1.add_run(company), bold=True)
    add_tab(p1)
    style_run(p1.add_run(dates), bold=True)

    # Role (Italic 11pt) ... Location (Regular 11pt, right-aligned)
    p2 = doc.add_paragraph()
    set_keep_with_next(p2)
    set_right_tabs(p2)
    set_spacing(p2, before=0, after=80, line=240)
    set_indent(p2, start=0)
    style_run(p2.add_run(role), italic=True)
    add_tab(p2)
    style_run(p2.add_run(location))
    return p1, p2


def add_project_header(doc, title, links, tech, first=False):
    # Title (Bold 11pt) ... Links (11pt, right-aligned).
    # First project after the section heading has before=0; later ones keep 280.
    p1 = doc.add_paragraph()
    set_keep_with_next(p1)
    set_right_tabs(p1)
    set_spacing(p1, before=0 if first else 280, after=0, line=240)
    set_indent(p1, start=0)
    style_run(p1.add_run(title), bold=True)
    add_tab(p1)
    for i, (name, url) in enumerate(links):
        if i > 0:
            style_run(p1.add_run(" | "))
        add_hyperlink(p1, name, url)

    # Tech stack (Italic 10pt)
    p2 = doc.add_paragraph()
    set_keep_with_next(p2)
    set_spacing(p2, before=0, after=80, line=240)
    set_indent(p2, start=0)
    style_run(p2.add_run(tech), size=TECH_SIZE, italic=True)
    return p1, p2


def add_cert_bullet(doc, name, url):
    """Certification row: hanging bullet (11pt inherited) + hyperlink."""
    p = doc.add_paragraph()
    set_spacing(p, before=0, after=40, line=240)
    set_indent(p, start=360, hanging=216)
    run = p.add_run()
    t_bullet = OxmlElement("w:t")
    t_bullet.text = "•"
    run._r.append(t_bullet)
    run._r.append(OxmlElement("w:tab"))
    add_hyperlink(p, name, url)
    return p


# ── Document Builder ──────────────────────────────────────────────────────────

def build_docx():
    doc = Document()
    set_document_defaults(doc)

    # A4 11906 x 16838 twips, 864-twip (0.6 in) margins — from CV (2).docx sectPr
    section = doc.sections[0]
    section.page_width = Twips(PAGE_W)
    section.page_height = Twips(PAGE_H)
    section.top_margin = Twips(MARGIN)
    section.bottom_margin = Twips(MARGIN)
    section.left_margin = Twips(MARGIN)
    section.right_margin = Twips(MARGIN)
    section.header_distance = Twips(0)
    section.footer_distance = Twips(0)

    # ── Header ────────────────────────────────────────────────────────────────
    p_name = doc.add_paragraph()
    p_name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_spacing(p_name, before=0, after=40, line=None)
    style_run(p_name.add_run("Jayesh Patel"), size=NAME_SIZE, bold=True)

    p_contact = doc.add_paragraph()
    p_contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_spacing(p_contact, before=0, after=120, line=None)
    style_run(p_contact.add_run("+91 8347205513  |  "))
    add_hyperlink(p_contact, "jayeshmepani777@gmail.com", "mailto:jayeshmepani777@gmail.com")
    style_run(p_contact.add_run("  |  "))
    add_hyperlink(p_contact, "jayeshmepani.site", "https://jayeshmepani.site/")
    style_run(p_contact.add_run("  |  "))
    add_hyperlink(p_contact, "linkedin.com/in/jayeshmepani", "https://linkedin.com/in/jayeshmepani/")
    style_run(p_contact.add_run("  |  "))
    add_hyperlink(p_contact, "github.com/jayeshmepani", "https://github.com/jayeshmepani/")

    # ── Professional Summary ──────────────────────────────────────────────────
    add_section_header(doc, "Professional Summary")
    p_sum = doc.add_paragraph()
    set_spacing(p_sum, before=0, after=120, line=240)
    set_indent(p_sum, start=0)
    style_run(p_sum.add_run(
        "Versatile Full-Stack & Systems Engineer with a B.Tech in Computer Science and deep expertise spanning "
        "systems-level programming, scalable backend architectures, high-performance computing, and cross-runtime "
        "integration. Proven track record designing and publishing production-grade open-source libraries, building "
        "high-throughput distributed APIs, and architecting data-intensive and intelligent software solutions. Strong "
        "foundation in software engineering fundamentals, dedicated to building performant, secure, and resilient "
        "systems from low-level runtimes to modern cloud and web platforms."
    ))

    # ── Technical Skills ──────────────────────────────────────────────────────
    add_section_header(doc, "Technical Skills")
    skills_data = [
        ("Languages", "C, C++, Rust, Python, PHP, JavaScript, TypeScript, Dart, Go, Java"),
        ("Compilers & Systems", "AST Parsers, FFI, Shared Libraries"),
        ("Frontend", "React, Next.js, Vue.js, AngularJS, Tailwind CSS, WebGL"),
        ("Backend", "Laravel, Node.js, Express.js, FastAPI, Flask, Django"),
        ("Databases", "PostgreSQL, MySQL, SQLite, MongoDB"),
        ("Mobile", "Flutter, Dart, Android (Java)"),
        ("Cloud & DevOps", "Linux, Nginx, GitLab CI, Render, Vercel"),
        ("AI & Data Science", "Deep Learning (PyTorch, TensorFlow), NLP, Computer Vision, Scikit-learn"),
        ("AI Media & Generative", "Stable Diffusion, GFPGAN, RealESRGAN, CodeFormer, UVR"),
        ("Design", "Figma, Adobe Creative Cloud"),
    ]
    for cat, val in skills_data:
        p = doc.add_paragraph()
        set_spacing(p, before=0, after=20, line=240)
        set_indent(p, start=0)
        style_run(p.add_run(f"{cat}: "), size=SKILL_SIZE, bold=True)
        style_run(p.add_run(val), size=SKILL_SIZE)

    # ── Work Experience ───────────────────────────────────────────────────────
    add_section_header(doc, "Work Experience")

    add_experience_header(doc, "Shreesoftech", "Dec 2024 – Present", "Laravel Developer", "Hybrid", first=True)
    add_bullet(doc, "Engineered production web architectures with Laravel, optimizing complex database workflows, Eloquent ORM queries, and RESTful API endpoints.")
    add_bullet(doc, "Implemented secure backend services, database migrations, and caching layers to ensure high availability and robust performance.")

    add_experience_header(doc, "WRTeam", "May 2024 – June 2024", "PHP Developer Intern", "Remote")
    add_bullet(doc, "Engineered and tested PHP web applications, improving codebase efficiency by 50%+ through dynamic fetching optimizations that significantly reduced file count and code volume.")
    add_bullet(doc, "Built secure API endpoints and collaborated on cross-functional agile teams to deliver commercial client solutions.")

    add_experience_header(doc, "CodSoft", "Jun 2024 – Jul 2024", "Flutter Developer Intern", "Remote")
    add_bullet(doc, "Built 4 cross-platform mobile applications with Flutter and Dart, implementing responsive UI components, state management, and local SQLite data persistence.")
    add_bullet(doc, "Conducted code reviews and optimized app performance across Android and iOS targets.")

    add_experience_header(doc, "CodSoft", "May 2024 – Jun 2024", "Python Programming Intern", "Remote")
    add_bullet(doc, "Constructed algorithmic automation suites, object-oriented software architectures, and backend scripting solutions across 5 Python applications spanning user management, game logic, security, and task management.")

    add_experience_header(doc, "CodSoft", "May 2024 – Jun 2024", "Data Science Intern", "Remote")
    add_bullet(doc, "Developed machine learning pipelines, predictive statistical modeling, and exploratory data analysis with Pandas, NumPy, and Scikit-learn across 5 datasets (55+ visualizations), achieving 85%+ model accuracy.")

    # ── Projects ──────────────────────────────────────────────────────────────
    add_section_header(doc, "Projects")

    add_project_header(
        doc,
        "Panchang Core & JPL Ephemeris Ecosystem",
        [
            ("GitHub (Core)", "https://github.com/jayeshmepani/panchang-core"),
            ("GitHub (C Engine)", "https://github.com/jayeshmepani/jpl-ephemeris"),
        ],
        "C, CALCEPH, JPL DE440, PHP ext-ffi, Python ctypes, Dart FFI",
        first=True,
    )
    add_bullet(doc, "Architected an authentic Vedic Panchanga calculation engine powered by an independent C ephemeris engine (204 public functions, 462 constants) implementing JPL DE405/DE440, CALCEPH, Moshier, VSOP87, and ELP2000 theories.")
    add_bullet(doc, "Computed Tithi, Nakṣatra, Yoga, Karaṇa, and 30 Muhūrtas with 0.001 arcsecond precision, orchestrating 336 unique festival identities and 126 vrat identities across Amanta and Purnimanta calendars.")
    add_bullet_with_parts(
        doc,
        [
            ("Distributed zero-overhead FFI bindings and precompiled native binaries across ", None),
            ("PHP", "https://github.com/jayeshmepani/jpl-moshier-ephemeris-php"),
            (" (ext-ffi), ", None),
            ("Python", "https://github.com/jayeshmepani/jpl-moshier-ephemeris-python"),
            (" (ctypes), and ", None),
            ("Dart/Flutter", "https://github.com/jayeshmepani/jpl-moshier-ephemeris-dart"),
            (" (dart:ffi).", None),
        ],
    )

    add_project_header(
        doc,
        "CSSForge — Lossless CSS AST Refactoring Engine",
        [
            ("GitHub", "https://github.com/jayeshmepani/cssforge"),
        ],
        "Rust (Edition 2024), Ratatui TUI, AST Parser, Crates.io",
    )
    add_bullet(doc, "Engineered a safety-first, lossless semantic CSS refactoring engine and interactive terminal workbench in Rust (Edition 2024), published on Crates.io.")
    add_bullet(doc, "Implemented 28 AST transformation rules modernizing flat legacy CSS into native nesting, Range media queries, @layer consolidation, and :is() factoring with zero declaration loss.")
    add_bullet(doc, "Supported raw stylesheets and embedded style blocks within Blade, Vue, Svelte, Astro, Twig, ERB, and HTML templates.")

    add_project_header(
        doc,
        "Lipimala — Deterministic Indic Script & Vedic Transliteration",
        [
            ("GitHub", "https://github.com/jayeshmepani/indic-script-converter"),
        ],
        "Dart, Node.js, Python, PHP, Unicode Vedic Tags",
    )
    add_bullet(doc, "Built a deterministic Indic transliteration suite published synchronously across 4 package registries: pub.dev (Dart), npm (JavaScript/TypeScript), PyPI (Python), and Packagist (PHP).")
    add_bullet(doc, "Solved Brahmic many-to-one collapse via structured TransliterationResult envelopes and checksummed LIT1: Unicode-Tag metadata trailers.")
    add_bullet(doc, "Enabled direct Devanagari ↔ Gujarati conversion with first-class preservation of Vedic svara markers (Udatta, Anudatta, Svarita).")

    add_project_header(
        doc,
        "Hindu Scriptures: Digital Repository & AI Scholar",
        [
            ("GitHub (Web)", "https://github.com/jayeshmepani/HinduScriptures/"),
            ("GitHub (Android)", "https://github.com/jayeshmepani/hindu-scripture-apk"),
        ],
        "Node.js, Express 5, Google Gemini AI, JavaScript, Java (Android APK)",
    )
    add_bullet(doc, "Developed a Sanātana Dharma digital repository indexing sacred literature across the Vedas, Upanishads, Bhagavad Gita, and Puranas with verse-level lexical search and Sanskrit root morphology.")
    add_bullet(doc, "Integrated contextual AI scripture analysis using Google Gemini and delivered both a responsive web platform and a native Android application.")

    add_project_header(
        doc,
        "Laravel Gemini AI Translation Extractor",
        [
            ("GitHub", "https://github.com/jayeshmepani/laravel-gemini-translator"),
        ],
        "PHP 8.3+, Laravel 11–13, Gemini AI, Fork (pcntl), Symfony Process, C kernel32 FFI",
    )
    add_bullet(doc, "Built an enterprise Laravel translation engine with dual-platform concurrency: Unix pcntl fork workers and Windows parallel Symfony Process workers, featuring smart token chunking, automatic rate-limit backoff, and atomic file writes.")
    add_bullet(doc, "Engineered a native Windows interactive TUI via C kernel32.dll FFI console hooks providing arrow-key navigation and multi-select prompts matching Laravel Prompts, alongside a web-based Translation Manager with script-fault detection across a 249-language catalog.")

    add_project_header(
        doc,
        "Poetry Analyzer — Deep Linguistic Analysis Engine",
        [("GitHub", "https://github.com/jayeshmepani/poetry-analyzer-app")],
        "Python, FastAPI, PyTorch, Transformers, spaCy, Stanza, HuggingFace, Alembic",
    )
    add_bullet(doc, "Designed an async FastAPI backend orchestrating 1.5 GB+ HuggingFace Transformer models, spaCy, Stanza, and Indic NLP to compute poetic meter, rhyme patterns, phonology, and semantic motifs with sub-second latency.")
    add_bullet(doc, "Implemented SQLAlchemy ORM connection pooling for SQLite, PostgreSQL, and MySQL, with Alembic schema migrations.")

    add_project_header(
        doc,
        "SQL to Laravel Migration Compiler",
        [
            ("GitHub", "https://github.com/jayeshmepani/sql-to-laravel"),
        ],
        "JavaScript, Web Workers, Topological Sort, Laravel 11/12/13",
    )
    add_bullet(doc, "Built a privacy-first, client-side compiler transforming raw SQL database dumps into idiomatic Laravel 11–13 migrations and seeders with topological foreign-key dependency resolution, vector/spatial type mapping, and environment-aware dialect generation across MySQL, PostgreSQL, MariaDB, and SQLite.")

    # ── Education ─────────────────────────────────────────────────────────────
    add_section_header(doc, "Education")

    p_inst = doc.add_paragraph()
    set_keep_with_next(p_inst)
    set_right_tabs(p_inst)
    set_spacing(p_inst, before=0, after=0, line=240)
    set_indent(p_inst, start=0)
    style_run(p_inst.add_run("Parul University"), bold=True)
    add_tab(p_inst)
    style_run(p_inst.add_run("Vadodara, Gujarat, India"))

    p_deg = doc.add_paragraph()
    set_keep_with_next(p_deg)
    set_right_tabs(p_deg)
    set_spacing(p_deg, before=0, after=0, line=240)
    set_indent(p_deg, start=0)
    style_run(p_deg.add_run("B.Tech in Computer Science and Engineering"), italic=True)
    add_tab(p_deg)
    style_run(p_deg.add_run("2021 – 2025"))

    p_cgpa = doc.add_paragraph()
    set_spacing(p_cgpa, before=0, after=40, line=None)
    set_indent(p_cgpa, start=0)
    style_run(p_cgpa.add_run("CGPA: 7.29"))

    # ── Certifications ────────────────────────────────────────────────────────
    add_section_header(doc, "Certifications")
    certs = [
        ("PHP Developer", "https://drive.google.com/file/d/1L1zCEEWwPwZR9s1vSDAtiGP-Eao2AlCV/view"),
        ("Python Programming", "https://drive.google.com/file/d/1L0sXZx5_cM2NbqPqMxu9LT66k9NIc714/view"),
        ("Data Science", "https://drive.google.com/file/d/1Kf9wefWCrOgYh_YBgHKEjPQgd8nGajpE/view"),
        ("Flutter Developer", "https://drive.google.com/file/d/1JvQcnijk_f-0-ApsC_HS0jfB2LpWnHAC/view"),
        ("Data Analytics with Python", "https://drive.google.com/file/d/1KsSZirAqiWB7n5yVZ7ywfdsdEItlR7dt/view"),
        ("Laravel", "https://drive.google.com/file/d/1h3f5_pZ8GIM2jUCLM_1oN1BcB0Qxi2Sg/view"),
    ]
    for name, url in certs:
        add_cert_bullet(doc, name, url)

    # ── Soft Skills ───────────────────────────────────────────────────────────
    add_section_header(doc, "Soft Skills")
    p_soft = doc.add_paragraph()
    set_spacing(p_soft, before=0, after=0, line=240)
    set_indent(p_soft, start=0)
    style_run(p_soft.add_run(
        "Analytical Thinking, Problem Solving, Systems Architecture, Attention to Detail, Technical Research, "
        "Time Management, Adaptability, Communication, Teamwork, Leadership"
    ))

    output_path = "cv_native.docx"
    doc.save(output_path)
    print(f"[SUCCESS] Native DOCX generated cleanly at: {os.path.abspath(output_path)}")


if __name__ == "__main__":
    build_docx()
